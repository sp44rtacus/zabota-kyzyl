import io
from datetime import datetime
from typing import List, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelService:
    """Сервис выгрузки отчетности АИС «Забота» в формате Microsoft Excel (.xlsx)"""

    @staticmethod
    def generate_applications_report(applications: List[Any], district_summary: List[dict]) -> bytes:
        """
        Генерация сводного аналитического отчета для руководства Департамента
        и Мэрии города Кызыла.
        """
        wb = openpyxl.Workbook()

        # Стили
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        bold_font = Font(name="Calibri", size=11, bold=True)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        sos_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        completed_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

        # --- Лист 1: Реестр обращений ---
        ws1 = wb.active
        ws1.title = "Реестр заявок"

        # Шапка документа
        ws1["A1"] = "МЭРИЯ ГОРОДА КЫЗЫЛА — ДЕПАРТАМЕНТ ПО СОЦИАЛЬНОЙ ПОЛИТИКЕ"
        ws1["A1"].font = title_font
        ws1["A2"] = f"Сводный реестр адресных заявок АИС «Забота» на {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ws1["A2"].font = Font(name="Calibri", size=11, italic=True, color="475569")

        headers = [
            "№", "УИН заявки", "Дата подачи", "ФИО заявителя", "Телефон", 
            "Категория льготы", "Микрорайон Кызыла", "Адрес", "Услуга", 
            "SOS (Экстренная)", "Текущий статус", "Исполнитель (Волонтер)"
        ]

        for col_idx, header in enumerate(headers, 1):
            cell = ws1.cell(row=4, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        for row_idx, app in enumerate(applications, 5):
            # Получение волонтера
            volunteer_str = "-"
            if app.assignments:
                first_assign = app.assignments[0]
                if first_assign.volunteer:
                    volunteer_str = first_assign.volunteer.full_name

            status_map = {
                "new": "Новая",
                "approved": "Одобрена",
                "in_progress": "В работе",
                "review": "На проверке",
                "completed": "Выполнена",
                "rejected": "Отклонена"
            }

            row_data = [
                row_idx - 4,
                app.uin,
                app.created_at.strftime("%d.%m.%Y %H:%M") if app.created_at else "-",
                app.applicant_name,
                app.phone,
                app.benefit_category,
                app.district.name if app.district else "-",
                app.address,
                app.category.title if app.category else "-",
                "ДА (SOS)" if app.is_emergency else "Нет",
                status_map.get(app.status, app.status),
                volunteer_str
            ]

            for col_idx, val in enumerate(row_data, 1):
                cell = ws1.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = left_align if col_idx in [4, 8, 9, 12] else center_align

                # Подсветка SOS и выполненных
                if app.is_emergency and col_idx == 10:
                    cell.fill = sos_fill
                    cell.font = Font(name="Calibri", size=11, bold=True, color="B91C1C")
                elif app.status == "completed" and col_idx == 11:
                    cell.fill = completed_fill
                    cell.font = Font(name="Calibri", size=11, bold=True, color="15803D")

        # Автоширина колонок для листа 1
        for col in ws1.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws1.column_dimensions[col_letter].width = max(max_len + 3, 11)

        # --- Лист 2: Сводная статистика по районам ---
        ws2 = wb.create_sheet(title="Статистика по районам")
        ws2["A1"] = "СВОДКА ОБРАЩЕНИЙ ПО МИКРОРАЙОНАМ ГОРОДА КЫЗЫЛА"
        ws2["A1"].font = title_font

        dist_headers = ["Микрорайон", "Всего заявок", "Выполнено", "Экстренные (SOS)", "Процент исполнения"]
        for col_idx, h in enumerate(dist_headers, 1):
            cell = ws2.cell(row=3, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        for r_idx, d in enumerate(district_summary, 4):
            total = d.get("total", 0)
            completed = d.get("completed", 0)
            rate = f"{(completed / total * 100):.1f}%" if total > 0 else "0.0%"

            row_vals = [
                d.get("name", "-"),
                total,
                completed,
                d.get("emergency", 0),
                rate
            ]
            for c_idx, val in enumerate(row_vals, 1):
                c = ws2.cell(row=r_idx, column=c_idx, value=val)
                c.border = thin_border
                c.alignment = left_align if c_idx == 1 else center_align
                if c_idx == 1:
                    c.font = bold_font

        for col in ws2.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws2.column_dimensions[col_letter].width = max(max_len + 4, 15)

        # Сохранение в байтовый буфер
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
